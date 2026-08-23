"""法规影响中心：监管事件影响评估查看/人工确认 + 一键创建批量复审批次。"""

import uuid
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, Query, Response, UploadFile
from sqlalchemy import func
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from src.api.middleware.auth import get_current_user, require_permission
from src.api.schemas.impact import (
    AssessmentBatchCreate,
    AssessmentConfirm,
    AssessmentCreate,
    AssessmentOut,
    MaterialCreate,
    MaterialOut,
    MaterialUpdate,
    RemediationDecision,
)
from src.api.schemas.rule import RuleOut
from src.core.impact import analyze
from src.core.impact.analysis import material_document_path
from src.core.impact.mining import mine_rule_candidates
from src.infrastructure.database.mysql import get_db
from src.models import (
    ImpactAssessment,
    PenaltyCase,
    ProductMaterial,
    RegulationUpdate,
    ReviewBatch,
    ReviewTask,
    Rule,
    RuleRegulationLink,
    User,
)
from src.services.impact_service import decide_remediation, start_retest
from src.services.review_service import start_review
from src.utils.config import settings
from src.utils.errors import AppError, ErrorCode
from src.utils.response import ok

router = APIRouter(prefix="/impact", tags=["impact"])

# 与 documents 路由的上传约束保持一致（解析器支持的类型）
RETEST_EXTS = {"pdf", "docx", "txt", "png", "jpg", "jpeg", "bmp"}
RETEST_MAX_SIZE = 50 * 1024 * 1024


def _get_assessment(db: Session, assessment_id: int) -> ImpactAssessment:
    assessment = db.get(ImpactAssessment, assessment_id)
    if assessment is None:
        raise AppError(ErrorCode.PARAM_ERROR, "影响评估不存在", http_status=404)
    return assessment


@router.get("/overview")
def impact_overview(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """风险驾驶舱：合规运营一屏指标（情报→评估→复审→整改→规则全链路计数）。"""
    pending_updates = (
        db.query(RegulationUpdate).filter(RegulationUpdate.status == "pending").count()
    )
    imported_ids = [
        row[0]
        for row in db.query(RegulationUpdate.id)
        .filter(RegulationUpdate.status == "imported")
        .all()
    ]
    analyzed_ids = {
        row[0]
        for row in db.query(ImpactAssessment.regulation_update_id)
        .filter(ImpactAssessment.regulation_update_id.isnot(None))
        .all()
    }
    to_analyze = len([rid for rid in imported_ids if rid not in analyzed_ids])

    status_rows = (
        db.query(ImpactAssessment.status, func.count(ImpactAssessment.id))
        .group_by(ImpactAssessment.status)
        .all()
    )
    assessments_by_status = {s: c for s, c in status_rows}
    pending_assessments = assessments_by_status.get("draft", 0) + assessments_by_status.get(
        "pending_review", 0
    )
    high_risk_materials = (
        db.query(func.sum(ImpactAssessment.high_count))
        .filter(ImpactAssessment.status.in_(("draft", "pending_review")))
        .scalar()
    ) or 0

    running_batches = (
        db.query(ReviewBatch)
        .filter(
            ReviewBatch.status == "running",
            ReviewBatch.source_type == "impact_assessment",
        )
        .count()
    )

    remediation_total = 0
    remediation_pending = 0
    remediation_done = 0
    for (items_json,) in db.query(ImpactAssessment.items).all():
        for entry in (items_json or {}).get("remediation") or []:
            remediation_total += 1
            if entry.get("decision") == "pending":
                remediation_pending += 1
            elif entry.get("decision") == "done":
                remediation_done += 1

    rules_pending_publish = (
        db.query(Rule)
        .filter(Rule.status.in_(("draft", "pending_review")))
        .count()
    )
    active_rules = db.query(Rule).filter(Rule.is_active == 1).count()

    return ok(
        {
            "pending_updates": pending_updates,
            "updates_to_analyze": to_analyze,
            "pending_assessments": pending_assessments,
            "assessments_by_status": assessments_by_status,
            "high_risk_materials": int(high_risk_materials),
            "running_batches": running_batches,
            "remediation_total": remediation_total,
            "remediation_pending": remediation_pending,
            "remediation_done": remediation_done,
            "rules_pending_publish": rules_pending_publish,
            "active_rules": active_rules,
        }
    )


@router.get("/assessments")
def list_assessments(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    status: str | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    query = db.query(ImpactAssessment).order_by(ImpactAssessment.created_at.desc())
    if status:
        query = query.filter(ImpactAssessment.status == status)
    total = query.count()
    rows = query.offset((page - 1) * page_size).limit(page_size).all()
    update_titles = {}
    for row in rows:
        if row.regulation_update_id and row.regulation_update_id not in update_titles:
            upd = db.get(RegulationUpdate, row.regulation_update_id)
            update_titles[row.regulation_update_id] = upd.title if upd else ""
    items = []
    for row in rows:
        data = AssessmentOut.model_validate(row).model_dump()
        data["update_title"] = update_titles.get(row.regulation_update_id, "")
        items.append(data)
    return ok({"total": total, "page": page, "page_size": page_size, "items": items})


@router.post("/assessments", status_code=201)
def create_assessment(
    body: AssessmentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("create_task")),
) -> dict:
    if (body.regulation_update_id is None) == (body.penalty_case_id is None):
        raise AppError(ErrorCode.PARAM_ERROR, "regulation_update_id 与 penalty_case_id 必须二选一")
    if body.regulation_update_id is not None:
        upd = db.get(RegulationUpdate, body.regulation_update_id)
        if upd is None:
            raise AppError(ErrorCode.REGULATION_UPDATE_NOT_FOUND, http_status=404)
    else:
        if db.get(PenaltyCase, body.penalty_case_id) is None:
            raise AppError(ErrorCode.PARAM_ERROR, "处罚案例不存在", http_status=404)
    assessment = analyze(
        db,
        regulation_update_id=body.regulation_update_id,
        penalty_case_id=body.penalty_case_id,
        created_by=current_user.id,
    )
    return ok(AssessmentOut.model_validate(assessment).model_dump())


@router.get("/materials")
def list_materials(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    product_type: str | None = Query(None),
    status: str | None = Query(None),
    last_review_status: str | None = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """产品材料台账：影响分析定位存量材料的底册。"""
    query = db.query(ProductMaterial).order_by(ProductMaterial.id.desc())
    if product_type:
        query = query.filter(ProductMaterial.product_type == product_type)
    if status:
        query = query.filter(ProductMaterial.status == status)
    if last_review_status:
        query = query.filter(ProductMaterial.last_review_status == last_review_status)
    total = query.count()
    rows = query.offset((page - 1) * page_size).limit(page_size).all()
    return ok(
        {
            "total": total,
            "page": page,
            "page_size": page_size,
            "items": [MaterialOut.model_validate(m).model_dump() for m in rows],
        }
    )


@router.post("/materials", status_code=201)
def create_material(
    body: MaterialCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("manage_rules")),
) -> dict:
    material = ProductMaterial(**body.model_dump())
    db.add(material)
    db.commit()
    db.refresh(material)
    return ok(MaterialOut.model_validate(material).model_dump())


@router.patch("/materials/{material_id}")
def update_material(
    material_id: int,
    body: MaterialUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("manage_rules")),
) -> dict:
    material = db.get(ProductMaterial, material_id)
    if material is None:
        raise AppError(ErrorCode.PARAM_ERROR, "材料不存在", http_status=404)
    updates = body.model_dump(exclude_unset=True)
    if not updates:
        raise AppError(ErrorCode.PARAM_ERROR, "没有需要更新的字段")
    for field, value in updates.items():
        setattr(material, field, value)
    db.commit()
    db.refresh(material)
    return ok(MaterialOut.model_validate(material).model_dump())


@router.post("/materials/{material_id}/retest", status_code=201)
async def retest_material(
    material_id: int,
    assessment_id: int = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("create_task")),
) -> dict:
    """整改复测：上传整改版材料（台账自动升版），单任务复审批次回挂原评估。

    复测通过（零违规）后原整改项自动标记 done、台账回写 clean；
    仍违规则新条目成为新的待整改项。
    """
    material = db.get(ProductMaterial, material_id)
    if material is None:
        raise AppError(ErrorCode.PARAM_ERROR, "材料不存在", http_status=404)
    assessment = _get_assessment(db, assessment_id)
    if assessment.status != "confirmed":
        raise AppError(ErrorCode.PARAM_ERROR, "整改复测须挂接已确认的影响评估")

    filename = file.filename or ""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in RETEST_EXTS:
        raise AppError(ErrorCode.FILE_TYPE_NOT_SUPPORTED)
    doc_id = "doc_" + uuid.uuid4().hex[:8]
    dest_dir = Path(settings.upload_dir) / str(current_user.id)
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / f"{doc_id}.{ext}"
    size = 0
    try:
        with dest.open("wb") as output:
            while chunk := await file.read(1024 * 1024):
                size += len(chunk)
                if size > RETEST_MAX_SIZE:
                    raise AppError(ErrorCode.FILE_TOO_LARGE)
                output.write(chunk)
    except Exception:
        dest.unlink(missing_ok=True)
        raise

    try:
        result = start_retest(db, assessment, material, str(dest), doc_id, current_user.id)
    except KeyError as exc:
        dest.unlink(missing_ok=True)
        raise AppError(ErrorCode.PARAM_ERROR, str(exc)) from None
    return ok({**result, "material_id": material.id, "document_id": doc_id})


@router.get("/penalty-cases")
def list_penalty_cases(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=50),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """处罚案例候选列表（影响中心触发源之二，倒序取最新）。"""
    total = db.query(PenaltyCase).count()
    rows = (
        db.query(PenaltyCase)
        .order_by(PenaltyCase.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    return ok(
        {
            "total": total,
            "items": [
                {
                    "id": c.id,
                    "case_name": c.case_name,
                    "penalty_org": c.penalty_org,
                    "penalty_date": c.penalty_date,
                    "violation_summary": c.violation_summary,
                    "related_regulation_count": len(c.related_regulations or []),
                }
                for c in rows
            ],
        }
    )


@router.post("/penalty-cases/{case_id}/mine-rules", status_code=201)
def mine_rules(
    case_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("manage_rules")),
) -> dict:
    """从罚单违规事实挖掘规则候选草稿（LLM 仅产草稿，人工发布后生效）。"""
    rules = mine_rule_candidates(db, case_id, created_by=current_user.id)
    return ok({"items": [RuleOut.model_validate(r).model_dump() for r in rules]})


@router.get("/assessments/{assessment_id}")
def get_assessment(
    assessment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    assessment = _get_assessment(db, assessment_id)
    data = AssessmentOut.model_validate(assessment).model_dump()
    upd = (
        db.get(RegulationUpdate, assessment.regulation_update_id)
        if assessment.regulation_update_id
        else None
    )
    data["update"] = (
        {"id": upd.id, "title": upd.title, "source": upd.source, "publish_date": upd.publish_date}
        if upd
        else None
    )
    batches = (
        db.query(ReviewBatch)
        .filter(
            ReviewBatch.source_type == "impact_assessment",
            ReviewBatch.source_id == assessment.id,
        )
        .order_by(ReviewBatch.id.desc())
        .all()
    )
    data["batches"] = [
        {
            "id": b.id,
            "batch_name": b.batch_name,
            "status": b.status,
            "total_count": b.total_count,
            "completed_count": b.completed_count,
        }
        for b in batches
    ]
    return ok(data)


@router.post("/assessments/{assessment_id}/confirm")
def confirm_assessment(
    assessment_id: int,
    body: AssessmentConfirm,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("confirm")),
) -> dict:
    """人工确认影响范围：确认后沉淀关键词命中的规则-法规关联，dismiss 直接作废。"""
    assessment = _get_assessment(db, assessment_id)
    if assessment.status not in ("draft", "pending_review"):
        raise AppError(
            ErrorCode.PARAM_ERROR, f"评估当前状态 {assessment.status}，不接受确认"
        )

    if body.action == "dismiss":
        assessment.status = "dismissed"
        db.commit()
        return ok(AssessmentOut.model_validate(assessment).model_dump())

    items = assessment.items or {}
    confirm_rule_ids = set(body.rule_ids)
    for entry in items.get("rules", []):
        if not confirm_rule_ids or entry.get("rule_id") in confirm_rule_ids:
            entry["confirm_status"] = "confirmed"
    confirm_material_ids = set(body.material_ids)
    for entry in items.get("materials", []):
        if not confirm_material_ids or entry.get("material_id") in confirm_material_ids:
            entry["confirm_status"] = "confirmed"
    # JSON 列原地改 dict 后 flush 按值比较判"无变化"，必须显式标脏
    flag_modified(assessment, "items")

    # 沉淀：把关键词兜底命中的规则-法规关联固化为结构化链（幂等：已有关联不重建）
    created_links = 0
    upd = (
        db.get(RegulationUpdate, assessment.regulation_update_id)
        if assessment.regulation_update_id
        else None
    )
    matched_reg_ids = [
        int(m["regulation_id"]) for m in (upd.matched_regulations or []) if m.get("regulation_id")
    ] if upd else []
    if matched_reg_ids:
        existing = {
            (link.rule_id, link.regulation_id)
            for link in db.query(RuleRegulationLink)
            .filter(RuleRegulationLink.regulation_id.in_(matched_reg_ids))
            .all()
        }
        for entry in items.get("rules", []):
            if entry.get("matched_by") != "keyword" or entry.get("confirm_status") != "confirmed":
                continue
            key = (entry.get("rule_id"), matched_reg_ids[0])
            if key in existing:
                continue
            db.add(
                RuleRegulationLink(
                    rule_id=entry["rule_id"],
                    regulation_id=matched_reg_ids[0],
                    citation=upd.title[:200],
                    source="impact_analysis",
                    created_by=current_user.id,
                )
            )
            existing.add(key)
            created_links += 1

    assessment.status = "confirmed"
    assessment.confirmed_by = current_user.id
    assessment.confirmed_at = datetime.now()
    db.commit()
    data = AssessmentOut.model_validate(assessment).model_dump()
    data["created_links"] = created_links
    return ok(data)


@router.post("/assessments/{assessment_id}/batch", status_code=201)
def create_review_batch(
    assessment_id: int,
    body: AssessmentBatchCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("create_task")),
) -> dict:
    """一键复审：默认圈选已确认的高/中风险材料，按台账 document_id 定位文档建批次。"""
    assessment = _get_assessment(db, assessment_id)
    if assessment.status != "confirmed":
        raise AppError(ErrorCode.PARAM_ERROR, "请先确认影响评估，再发起批量复审")

    wanted = set(body.material_ids)
    selected = []
    for entry in (assessment.items or {}).get("materials", []):
        if wanted:
            if entry.get("material_id") not in wanted:
                continue
        elif not (
            entry.get("risk_level") in ("high", "medium")
            and entry.get("confirm_status") == "confirmed"
        ):
            continue
        selected.append(entry)
    if not selected:
        raise AppError(ErrorCode.PARAM_ERROR, "没有符合条件的可复审材料")

    skipped: list[dict] = []
    resolved: list[tuple[ProductMaterial, Path]] = []
    for entry in selected:
        material = db.get(ProductMaterial, entry["material_id"])
        if material is None:
            skipped.append({"material_id": entry["material_id"], "reason": "台账记录不存在"})
            continue
        path = material_document_path(material.document_id)
        if path is None:
            skipped.append(
                {"material_id": material.id, "name": material.name, "reason": "未关联可定位的文档"}
            )
            continue
        resolved.append((material, path))
    if not resolved:
        raise AppError(
            ErrorCode.PARAM_ERROR, f"所选材料均无文档可复审（{len(skipped)} 份被跳过）"
        )

    batch = ReviewBatch(
        batch_name=f"复审批次：{assessment.title}"[:255],
        total_count=len(resolved),
        created_by=current_user.id,
        source_type="impact_assessment",
        source_id=assessment.id,
    )
    db.add(batch)
    db.flush()
    task_ids = []
    for material, path in resolved:
        task = ReviewTask(
            task_name=material.name,
            document_name=path.name,
            document_path=str(path),
            status="pending",
            created_by=current_user.id,
            batch_id=batch.id,
        )
        db.add(task)
        db.flush()
        task_ids.append(task.id)
    db.commit()

    for task_id in task_ids:
        start_review(task_id)
    return ok(
        {
            "batch_id": batch.id,
            "batch_name": batch.batch_name,
            "total_count": batch.total_count,
            "task_ids": task_ids,
            "skipped": skipped,
        }
    )


@router.get("/assessments/{assessment_id}/remediation")
def get_remediation(
    assessment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> dict:
    """整改清单：复审批次终态后自动生成，含违规计数与人工决策。"""
    assessment = _get_assessment(db, assessment_id)
    return ok({"items": (assessment.items or {}).get("remediation") or []})


@router.get("/assessments/{assessment_id}/remediation/export")
def export_remediation(
    assessment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("view_report")),
) -> Response:
    """整改清单导出 Excel（合规向管理层汇报用）。"""
    from io import BytesIO
    from urllib.parse import quote

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    assessment = _get_assessment(db, assessment_id)
    items = (assessment.items or {}).get("remediation") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "整改清单"
    headers = ["序号", "材料", "违规数", "需整改", "决策", "复测", "关联任务", "任务状态", "批次"]
    ws.append([f"整改清单 —— {assessment.title}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(headers)
    for cell in ws[2]:
        cell.font = Font(bold=True)

    decision_labels = {
        "pending": "待决策",
        "remediate": "整改中",
        "reject": "已驳回",
        "done": "已完成",
    }
    for i, e in enumerate(items, start=1):
        ws.append(
            [
                i,
                e.get("material_name", ""),
                e.get("violation_count", 0),
                "是" if e.get("requires_remediation") else "否",
                decision_labels.get(e.get("decision"), e.get("decision", "")),
                "复测条目" if e.get("retest") else "",
                f"#{e.get('task_id', '')}",
                e.get("task_status", ""),
                f"#{e.get('batch_id', '')}",
            ]
        )
    widths = [6, 36, 8, 8, 10, 10, 12, 12, 10]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = quote(f"整改清单-{assessment.id}.xlsx")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.post("/assessments/{assessment_id}/remediation")
def set_remediation_decision(
    assessment_id: int,
    body: RemediationDecision,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("confirm")),
) -> dict:
    try:
        items = decide_remediation(db, assessment_id, body.task_id, body.decision)
    except KeyError as exc:
        raise AppError(ErrorCode.PARAM_ERROR, str(exc), http_status=404) from exc
    return ok({"items": items})
