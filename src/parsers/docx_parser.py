"""Word 解析器：基于 python-docx 提取段落，识别标题样式。
python-docx 无分页信息，以段落序号作为位置锚点。"""

from pathlib import Path

from docx import Document as DocxDocument

from src.parsers.base import DocumentParser, ParsedDocument, TextBlock


class DocxParser(DocumentParser):
    def parse(self, path: str | Path) -> ParsedDocument:
        file_path = Path(path)
        doc = DocxDocument(str(file_path))
        blocks: list[TextBlock] = []
        for idx, para in enumerate(doc.paragraphs, start=1):
            text = para.text.strip()
            if not text:
                continue
            is_heading = para.style.name.lower().startswith("heading")
            blocks.append(
                TextBlock(text=text, page=None, paragraph=idx, heading=is_heading)
            )
        return ParsedDocument(
            file_name=file_path.name,
            title=blocks[0].text if blocks else "",
            blocks=blocks,
        )
