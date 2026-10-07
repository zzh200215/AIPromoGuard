"""Excel 解析器：基于 openpyxl，逐行展平为文本块。

合规审查关注单元格文案（收益率、费率、承诺类表述常藏在表格里）；
每个工作表标题行标记 heading，行内单元格以 " | " 连接，空行跳过。
read_only + data_only：取公式计算值，避免大文件整表载入内存。
"""

from pathlib import Path

import openpyxl

from src.parsers.base import DocumentParser, ParsedDocument, TextBlock

_MAX_ROWS_PER_SHEET = 2000
_MAX_COLS = 30


class XlsxParser(DocumentParser):
    def parse(self, path: str | Path) -> ParsedDocument:
        file_path = Path(path)
        wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
        blocks: list[TextBlock] = []
        try:
            for sheet_no, ws in enumerate(wb.worksheets, start=1):
                para_no = 0
                blocks.append(
                    TextBlock(
                        text=f"[工作表] {ws.title.strip()}",
                        page=sheet_no,
                        paragraph=para_no + 1,
                        heading=True,
                    )
                )
                para_no += 1
                for row in ws.iter_rows(
                    max_row=_MAX_ROWS_PER_SHEET, max_col=_MAX_COLS, values_only=True
                ):
                    cells = [
                        str(c).strip() for c in row if c is not None and str(c).strip()
                    ]
                    if not cells:
                        continue
                    para_no += 1
                    blocks.append(
                        TextBlock(text=" | ".join(cells), page=sheet_no, paragraph=para_no)
                    )
        finally:
            wb.close()
        return ParsedDocument(
            file_name=file_path.name,
            title=blocks[0].text if blocks else "",
            blocks=blocks,
        )
