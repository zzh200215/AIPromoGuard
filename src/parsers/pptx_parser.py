"""PPT 解析器：基于 python-pptx，逐张幻灯片提取文本框/表格/备注。

营销 PPT 的合规审查主要看页面文案与表格；每张幻灯片作为一"页"，
标题占位符标记 heading（分块/断句优先按标题切）。
"""

from pathlib import Path

from pptx import Presentation

from src.parsers.base import DocumentParser, ParsedDocument, TextBlock

_MAX_SLIDES = 200


class PptxParser(DocumentParser):
    def parse(self, path: str | Path) -> ParsedDocument:
        file_path = Path(path)
        prs = Presentation(str(file_path))
        blocks: list[TextBlock] = []
        for slide_no, slide in enumerate(prs.slides, start=1):
            if slide_no > _MAX_SLIDES:
                break
            para_no = 0
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.has_text_frame:
                text = title_shape.text_frame.text.strip()
                if text:
                    para_no += 1
                    blocks.append(
                        TextBlock(text=text, page=slide_no, paragraph=para_no, heading=True)
                    )
            for shape in slide.shapes:
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        line = " | ".join(c for c in cells if c).strip(" |")
                        if line:
                            para_no += 1
                            blocks.append(
                                TextBlock(text=line, page=slide_no, paragraph=para_no)
                            )
                    continue
                if not shape.has_text_frame:
                    continue
                for raw in shape.text_frame.text.splitlines():
                    text = raw.strip()
                    if not text:
                        continue
                    para_no += 1
                    blocks.append(TextBlock(text=text, page=slide_no, paragraph=para_no))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    para_no += 1
                    blocks.append(
                        TextBlock(text=f"[演讲备注] {notes}", page=slide_no, paragraph=para_no)
                    )
        return ParsedDocument(
            file_name=file_path.name,
            title=blocks[0].text if blocks else "",
            blocks=blocks,
        )
