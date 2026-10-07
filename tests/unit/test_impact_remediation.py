"""闭环收口测试：批次回写台账 + 整改清单人工决策 + 规则版本发布。"""

from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from scripts.demo.seed_impact_demo import _demo_dir, reset_demo, seed_demo
from sqlalchemy import text
from src.api.main import app
from src.core.impact import analyze
from src.infrastructure.database.mysql import SessionLocal
from src.models import (
    ImpactAssessment,
    ProductMaterial,
    ReviewBatch,
    ReviewResult,
    ReviewTask,
    Rule,
)
from src.services.impact_service import sync_batch_to_assessment

client = TestClient(app)


def _db_ready() -> bool:
    try:
        db = SessionLocal()
        db.execute(text("SELECT 1"))
        db.close()
        return True
    except Exception:
        return False


def _admin_headers() -> dict:
    login = client.post("/api/v1/auth/login", json={"username": "admin", "password": "admin123"})
    return {"Authorization": f"Bearer {login.json()['data']['access_token']}"}


def _seed_review_batch(db, assessment_id: int, doc_ids: list[str]) -> tuple[ReviewBatch, list[int]]:
    batch = ReviewBatch(
        batch_name="回写测试批次",
        status="completed",
        total_count=len(doc_ids),
        completed_count=len(doc_ids),
        source_type="impact_assessment",
        source_id=assessment_id,
    )
    db.add(batch)
    db.flush()
    task_ids = []
    for doc_id in doc_ids:
        task = ReviewTask(
            task_name=f"回写-{doc_id}",
            document_name=f"{doc_id}.txt",
            document_path=str(_demo_dir() / f"{doc_id}.txt"),
            status="completed",
            batch_id=batch.id,
        )
        db.add(task)
        db.flush()
        task_ids.append(task.id)
    db.commit()
    return batch, task_ids


@pytest.mark.skipif(not _db_ready(), reason="数据库不可用")
def test_sync_batch_and_remediation_decision(isolated_materials):
    db = SessionLocal()
    try:
        ids = seed_demo(db)
        assessment = analyze(db, ids["regulation_update_id"])
        doc_ids = ["doc_5d6f7a01", "doc_5d6f7a07"]  # 一份高违规材料 + 一份干净材料
        batch, task_ids = _seed_review_batch(db, assessment.id, doc_ids)

        # 材料一有 1 条违规命中，材料二干净
        db.add(
            ReviewResult(
                task_id=task_ids[0],
                violation_type="rule",
                risk_level="high",
                original_text="保本无虞",
                confidence_score=Decimal("0.9"),
            )
        )
        db.commit()

        sync_batch_to_assessment(db, batch.id)

        remediation = (db.get(ImpactAssessment, assessment.id).items or {})["remediation"]
        assert len(remediation) == 2
        by_task = {e["task_id"]: e for e in remediation}
        assert by_task[task_ids[0]]["requires_remediation"] is True
        assert by_task[task_ids[1]]["requires_remediation"] is False
        db.expire_all()
        statuses = {
            m.document_id: m.last_review_status
            for m in db.query(ProductMaterial)
            .filter(ProductMaterial.document_id.in_(doc_ids))
            .all()
        }
        assert statuses[doc_ids[0]] == "violation"
        assert statuses[doc_ids[1]] == "clean"

        # 人工决策落库且重放同步不覆盖
        h = _admin_headers()
        r = client.post(
            f"/api/v1/impact/assessments/{assessment.id}/remediation",
            headers=h,
            json={"task_id": task_ids[0], "decision": "remediate"},
        )
        assert r.status_code == 200, r.text
        assert any(
            e["decision"] == "remediate" for e in r.json()["data"]["items"]
        )

        db.rollback()  # 结束本会话读快照（REPEATABLE READ），重读 API 会话的决策
        sync_batch_to_assessment(db, batch.id)  # 幂等重放
        items = (db.get(ImpactAssessment, assessment.id).items or {})["remediation"]
        assert len(items) == 2
        assert next(e for e in items if e["task_id"] == task_ids[0])["decision"] == "remediate"
    finally:
        reset_demo(db)
        db.close()


@pytest.mark.skipif(not _db_ready(), reason="数据库不可用")
def test_remediation_export_excel(isolated_materials):
    from io import BytesIO
    from urllib.parse import unquote

    from openpyxl import load_workbook
    from sqlalchemy.orm.attributes import flag_modified

    db = SessionLocal()
    try:
        ids = seed_demo(db)
        assessment = analyze(db, ids["regulation_update_id"])
        items = assessment.items or {}
        items["remediation"] = [
            {
                "batch_id": 1,
                "task_id": 1,
                "task_status": "completed",
                "material_id": None,
                "material_name": "导出测试材料",
                "violation_count": 2,
                "requires_remediation": True,
                "decision": "remediate",
            }
        ]
        assessment.items = items
        flag_modified(assessment, "items")
        db.commit()

        h = _admin_headers()
        r = client.get(
            f"/api/v1/impact/assessments/{assessment.id}/remediation/export", headers=h
        )
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
        )
        disposition = unquote(r.headers["content-disposition"])  # filename* 为 UTF-8 百分号编码
        assert "整改清单" in disposition and ".xlsx" in disposition
        wb = load_workbook(BytesIO(r.content))
        ws = wb.active
        assert ws["A1"].value and str(ws["A1"].value).startswith("整改清单")
        rows = list(ws.iter_rows(values_only=True))
        assert any("导出测试材料" in str(row[1] or "") for row in rows)
        assert any("整改中" in str(row[4] or "") for row in rows)
    finally:
        reset_demo(db)
        db.close()


@pytest.mark.skipif(not _db_ready(), reason="数据库不可用")
def test_rule_version_publish_flow():
    db = SessionLocal()
    h = _admin_headers()
    rule_name = "发布流程测试规则"
    try:
        created = client.post(
            "/api/v1/rules",
            headers=h,
            json={
                "domain": "fund",
                "rule_name": rule_name,
                "rule_type": "regex",
                "rule_content": {
                    "operator": "OR",
                    "conditions": [{"type": "regex", "pattern": "测试甲"}],
                },
                "severity": "low",
            },
        )
        assert created.status_code == 201, created.text
        v1 = created.json()["data"]
        # 规则发布门禁（R4-2）：新建规则默认 draft（is_active=0），
        # 需显式发布才生效——原断言 status == "active" 写于门禁功能之前，已过时
        assert v1["status"] == "draft" and v1["is_active"] == 0

        # 发布门禁流程：v1 先发布生效——否则"取代"语义不成立（draft 版本从未
        # 生效，发布新版本时它不在 superseded 列表里，这是正确实现语义）
        v1_pub = client.post(f"/api/v1/rules/{v1['id']}/publish", headers=h)
        assert v1_pub.status_code == 200, v1_pub.text
        assert v1_pub.json()["data"]["rule"]["status"] == "active"

        draft = client.post(f"/api/v1/rules/{v1['id']}/versions", headers=h).json()["data"]
        assert draft["status"] == "draft" and draft["version"] == v1["version"] + 1
        assert draft["is_active"] == 0

        published = client.post(f"/api/v1/rules/{draft['id']}/publish", headers=h)
        assert published.status_code == 200, published.text
        data = published.json()["data"]
        assert data["rule"]["status"] == "active" and data["rule"]["is_active"] == 1
        assert v1["id"] in data["superseded_rule_ids"]
        old = db.get(Rule, v1["id"])
        assert old.is_active == 0  # 旧版本自动停用

        again = client.post(f"/api/v1/rules/{draft['id']}/publish", headers=h)
        assert again.status_code == 400  # 已生效版本不能再发布
    finally:
        db.query(Rule).filter(Rule.rule_name == rule_name).delete(synchronize_session=False)
        db.commit()
        db.close()
