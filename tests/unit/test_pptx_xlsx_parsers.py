"""PPT / Excel 解析器测试：用 python-pptx / openpyxl 现场生成真实文件再解析。

营销合规审查的高频场景（PPT 文案、Excel 费率表）此前不支持，本组测试锁定
提取行为：幻灯片标题→heading、表格行、备注；工作表标题→heading、单元格行。
"""

from src.parsers.pptx_parser import PptxParser
from src.parsers.xlsx_parser import XlsxParser


def _make_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    # 第一页：标题 + 正文 + 表格 + 备注
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "稳健回报基金产品介绍"
    body = slide.placeholders[1]
    body.text = "本金安全无忧，稳赚不赔\n历史年化收益超 8%"
    rows, cols = 2, 2
    table = slide.shapes.add_table(rows, cols, Inches(1), Inches(3), Inches(8), Inches(1)).table
    table.cell(0, 0).text = "费率类型"
    table.cell(0, 1).text = "费率"
    table.cell(1, 0).text = "申购费"
    table.cell(1, 1).text = "0.8%"
    slide.notes_slide.notes_text_frame.text = "内部提醒：收益率表述需合规确认"
    path = tmp_path / "promo.pptx"
    prs.save(str(path))
    return path


def test_pptx_parser_extracts_title_body_table_notes(tmp_path):
    path = _make_pptx(tmp_path)
    doc = PptxParser().parse(str(path))
    texts = [b.text for b in doc.blocks]
    assert "稳健回报基金产品介绍" in texts
    assert doc.blocks[0].heading is True and doc.blocks[0].page == 1
    assert "本金安全无忧，稳赚不赔" in texts
    assert "申购费 | 0.8%" in texts
    assert any(b.text.startswith("[演讲备注]") and "合规确认" in b.text for b in doc.blocks)


def test_pptx_parser_empty_deck(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    path = tmp_path / "empty.pptx"
    prs.save(str(path))
    doc = PptxParser().parse(str(path))
    assert doc.blocks == []


def _make_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "费率表"
    ws.append(["产品", "申购费", "赎回费"])
    ws.append(["稳健回报", 0.008, 0.005])
    ws.append(["稳赚不赔专享", 0, 0])
    ws2 = wb.create_sheet("宣传语")
    ws2.append(["本金安全无忧"])
    path = tmp_path / "rates.xlsx"
    wb.save(str(path))
    return path


def test_xlsx_parser_flattens_rows_with_sheet_headings(tmp_path):
    path = _make_xlsx(tmp_path)
    doc = XlsxParser().parse(str(path))
    texts = [b.text for b in doc.blocks]
    assert texts[0] == "[工作表] 费率表" and doc.blocks[0].heading is True
    assert "产品 | 申购费 | 赎回费" in texts
    assert "稳健回报 | 0.008 | 0.005" in texts
    assert "稳赚不赔专享 | 0 | 0" in texts
    assert "[工作表] 宣传语" in texts and "本金安全无忧" in texts
    # 每个工作表的 page 序号递增
    pages = [b.page for b in doc.blocks if b.heading]
    assert pages == [1, 2]


def test_xlsx_parser_skips_blank_rows(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "s"
    ws["A1"] = "有内容"
    ws["A2"] = None
    ws["A3"] = "  "  # 纯空白
    ws["A4"] = "末行"
    path = tmp_path / "blank.xlsx"
    wb.save(str(path))
    doc = XlsxParser().parse(str(path))
    texts = [b.text for b in doc.blocks if not b.heading]
    assert texts == ["有内容", "末行"]
