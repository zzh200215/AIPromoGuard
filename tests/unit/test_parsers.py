from pathlib import Path

import pytest
from src.parsers import parse_document
from src.utils.errors import AppError, ErrorCode

_FONT_CANDIDATES = [
    "C:/Windows/Fonts/simhei.ttf",
    "C:/Windows/Fonts/simsun.ttc",
    "C:/Windows/Fonts/msyh.ttc",
]


def _find_cjk_font():
    pytest.importorskip("fpdf")  # PyPI 包名 fpdf2，模块名为 fpdf
    for font in _FONT_CANDIDATES:
        if Path(font).exists():
            return font
    pytest.skip("未找到 CJK 字体，跳过 PDF 测试")


def _write_pdf(path, lines):
    from fpdf import FPDF

    font = _find_cjk_font()
    pdf = FPDF()
    pdf.add_font("cjk", "", font)
    for line in lines:
        pdf.add_page()
        pdf.set_font("cjk", "", 16)
        pdf.cell(0, 10, line, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(path))


def test_txt_parser(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("第一行\n\n第二行\n\n第三行", encoding="utf-8")
    doc = parse_document(str(p))
    assert len(doc.blocks) == 3
    assert [b.paragraph for b in doc.blocks] == [1, 3, 5]
    assert all(b.page == 1 for b in doc.blocks)


def test_txt_gbk_encoding(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes("中文测试".encode("gbk"))
    doc = parse_document(str(p))
    assert doc.blocks[0].text == "中文测试"


def test_docx_parser(tmp_path):
    from docx import Document

    p = tmp_path / "a.docx"
    d = Document()
    d.add_heading("说明书标题", level=1)
    d.add_paragraph("正文第一段")
    d.add_paragraph("正文第二段")
    d.save(str(p))

    doc = parse_document(str(p))
    assert len(doc.blocks) == 3
    assert doc.blocks[0].heading is True
    assert doc.blocks[1].heading is False
    assert doc.blocks[1].paragraph == 2


def test_pdf_parser_multipage(tmp_path):
    p = tmp_path / "a.pdf"
    _write_pdf(p, ["第一页内容", "第二页内容", "第三页内容"])
    doc = parse_document(str(p))
    assert [b.page for b in doc.blocks] == [1, 2, 3]
    assert doc.blocks[0].text == "第一页内容"
    assert doc.blocks[2].text == "第三页内容"


def test_unsupported_format(tmp_path):
    p = tmp_path / "a.exe"
    p.write_bytes(b"MZ")
    with pytest.raises(AppError) as exc_info:
        parse_document(str(p))
    assert exc_info.value.code == ErrorCode.FILE_TYPE_NOT_SUPPORTED
